"""
Excel样式美化服务
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo
from typing import Dict, Any
import pandas as pd


class ExcelStyleService:
    """Excel样式美化服务"""
    
    # 颜色定义
    HEADER_COLOR = "2E86AB"  # 深蓝色
    HEADER_TEXT_COLOR = "FFFFFF"  # 白色
    SECONDARY_COLOR = "A23B72"  # 紫色
    SUCCESS_COLOR = "F18F01"  # 橙色
    WARNING_COLOR = "C73E1D"  # 红色
    LIGHT_GRAY = "F8F9FA"  # 浅灰色
    BORDER_COLOR = "DEE2E6"  # 边框色
    
    @staticmethod
    def apply_header_style(worksheet, start_row=1, end_col=None):
        """应用表头样式"""
        if end_col is None:
            end_col = worksheet.max_column
            
        # 定义样式
        header_font = Font(
            name='Microsoft YaHei',
            size=12,
            bold=True,
            color=ExcelStyleService.HEADER_TEXT_COLOR
        )
        
        header_fill = PatternFill(
            start_color=ExcelStyleService.HEADER_COLOR,
            end_color=ExcelStyleService.HEADER_COLOR,
            fill_type='solid'
        )
        
        header_alignment = Alignment(
            horizontal='center',
            vertical='center',
            wrap_text=True
        )
        
        # 应用样式
        for col in range(1, end_col + 1):
            cell = worksheet.cell(row=start_row, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
    
    @staticmethod
    def apply_data_style(worksheet, start_row=2, end_row=None, end_col=None):
        """应用数据行样式"""
        if end_row is None:
            end_row = worksheet.max_row
        if end_col is None:
            end_col = worksheet.max_column
            
        # 定义样式
        data_font = Font(name='Microsoft YaHei', size=11)
        data_alignment = Alignment(
            horizontal='left',
            vertical='center',
            wrap_text=True
        )
        
        # 边框样式
        thin_border = Border(
            left=Side(style='thin', color=ExcelStyleService.BORDER_COLOR),
            right=Side(style='thin', color=ExcelStyleService.BORDER_COLOR),
            top=Side(style='thin', color=ExcelStyleService.BORDER_COLOR),
            bottom=Side(style='thin', color=ExcelStyleService.BORDER_COLOR)
        )
        
        # 应用样式
        for row in range(start_row, end_row + 1):
            is_even = (row - start_row) % 2 == 0
            for col in range(1, end_col + 1):
                cell = worksheet.cell(row=row, column=col)
                cell.font = data_font
                cell.alignment = data_alignment
                cell.border = thin_border
                
                # 交替行颜色
                if is_even:
                    cell.fill = PatternFill(
                        start_color=ExcelStyleService.LIGHT_GRAY,
                        end_color=ExcelStyleService.LIGHT_GRAY,
                        fill_type='solid'
                    )
    
    @staticmethod
    def apply_column_widths(worksheet, column_widths=None):
        """设置列宽"""
        if column_widths is None:
            # 默认列宽
            for col in range(1, worksheet.max_column + 1):
                worksheet.column_dimensions[worksheet.cell(row=1, column=col).column_letter].width = 15
        else:
            for col, width in column_widths.items():
                if isinstance(col, int):
                    col = worksheet.cell(row=1, column=col).column_letter
                worksheet.column_dimensions[col].width = width
    
    @staticmethod
    def add_title(worksheet, title, subtitle=None):
        """添加标题"""
        # 标题样式
        title_font = Font(
            name='Microsoft YaHei',
            size=16,
            bold=True,
            color=ExcelStyleService.HEADER_COLOR
        )
        
        subtitle_font = Font(
            name='Microsoft YaHei',
            size=12,
            color=ExcelStyleService.SECONDARY_COLOR
        )
        
        # 插入标题行
        worksheet.insert_rows(1)
        worksheet.insert_rows(2)
        
        # 设置标题
        worksheet.cell(row=1, column=1, value=title)
        worksheet.cell(row=1, column=1).font = title_font
        
        if subtitle:
            worksheet.cell(row=2, column=1, value=subtitle)
            worksheet.cell(row=2, column=1).font = subtitle_font
        
        return 3  # 返回数据开始的行号
    
    @staticmethod
    def add_summary_info(worksheet, summary_data, start_row=None):
        """添加汇总信息"""
        if start_row is None:
            start_row = worksheet.max_row + 2
            
        # 汇总样式
        summary_font = Font(
            name='Microsoft YaHei',
            size=11,
            bold=True
        )
        
        summary_fill = PatternFill(
            start_color=ExcelStyleService.SUCCESS_COLOR,
            end_color=ExcelStyleService.SUCCESS_COLOR,
            fill_type='solid'
        )
        
        # 添加汇总信息
        for i, (key, value) in enumerate(summary_data.items(), 1):
            worksheet.cell(row=start_row + i - 1, column=1, value=f"{key}: {value}")
            cell = worksheet.cell(row=start_row + i - 1, column=1)
            cell.font = summary_font
            cell.fill = summary_fill
    
    @staticmethod
    def create_styled_worksheet(df, sheet_name, title=None, subtitle=None, 
                              column_widths=None, summary_data=None):
        """创建带样式的Excel工作表"""
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        
        # 添加标题
        data_start_row = 1
        if title:
            data_start_row = ExcelStyleService.add_title(ws, title, subtitle)
        
        # 写入数据
        for r in dataframe_to_rows(df, index=False, header=True):
            ws.append(r)
        
        # 应用样式
        header_row = data_start_row
        data_start_row += 1
        ExcelStyleService.apply_header_style(ws, header_row)
        ExcelStyleService.apply_data_style(ws, data_start_row)
        
        # 设置列宽
        ExcelStyleService.apply_column_widths(ws, column_widths)
        
        # 添加汇总信息
        if summary_data:
            ExcelStyleService.add_summary_info(ws, summary_data)
        
        return wb, ws
    
    @staticmethod
    def create_user_export_excel(df):
        """创建用户导出Excel"""
        # 列宽设置
        column_widths = {
            'A': 8,   # ID
            'B': 15,  # 用户名
            'C': 25,  # 邮箱
            'D': 15,  # 真实姓名
            'E': 10,  # 角色
            'F': 10,  # 状态
            'G': 20   # 创建时间
        }
        
        # 汇总数据
        total_users = len(df)
        admin_count = len(df[df['角色'] == 'admin']) if total_users > 0 else 0
        user_count = len(df[df['角色'] == 'user']) if total_users > 0 else 0
        active_count = len(df[df['状态'] == 'active']) if total_users > 0 else 0
        
        summary_data = {
            '总用户数': total_users,
            '管理员数': admin_count,
            '普通用户数': user_count,
            '活跃用户数': active_count
        }
        
        wb, ws = ExcelStyleService.create_styled_worksheet(
            df=df,
            sheet_name='用户数据',
            title='ExamSphere 用户数据导出',
            subtitle=f'导出时间: {pd.Timestamp.now().strftime("%Y-%m-%d %H:%M:%S")}',
            column_widths=column_widths,
            summary_data=summary_data
        )
        
        return wb
    
    @staticmethod
    def create_subject_export_excel(df):
        """创建科目导出Excel"""
        # 列宽设置
        column_widths = {
            'A': 8,   # ID
            'B': 20,  # 科目名称
            'C': 15,  # 科目代码
            'D': 30,  # 描述
            'E': 15,  # 分类
            'F': 10,  # 是否免费
            'G': 10,  # 价格
            'H': 10,  # 状态
            'I': 20   # 创建时间
        }
        
        # 汇总数据
        total_subjects = len(df)
        active_count = len(df[df['状态'] == 'active']) if total_subjects > 0 else 0
        free_count = len(df[df['是否免费'] == '是']) if total_subjects > 0 else 0
        
        summary_data = {
            '总科目数': total_subjects,
            '活跃科目数': active_count,
            '免费科目数': free_count,
            '收费科目数': total_subjects - free_count
        }
        
        wb, ws = ExcelStyleService.create_styled_worksheet(
            df=df,
            sheet_name='科目数据',
            title='ExamSphere 科目数据导出',
            subtitle=f'导出时间: {pd.Timestamp.now().strftime("%Y-%m-%d %H:%M:%S")}',
            column_widths=column_widths,
            summary_data=summary_data
        )
        
        return wb
    
    @staticmethod
    def create_question_export_excel(df):
        """创建试题导出Excel"""
        # 列宽设置
        column_widths = {
            'A': 8,   # ID
            'B': 8,   # 科目ID
            'C': 10,  # 题型
            'D': 40,  # 题目
            'E': 30,  # 内容
            'F': 30,  # 选项
            'G': 20,  # 答案
            'H': 30,  # 解析
            'I': 8,   # 难度
            'J': 8,   # 分值
            'K': 10,  # 状态
            'L': 20   # 创建时间
        }
        
        # 汇总数据
        total_questions = len(df)
        if total_questions > 0:
            single_count = len(df[df['题型'] == 'single'])
            multiple_count = len(df[df['题型'] == 'multiple'])
            judge_count = len(df[df['题型'] == 'judge'])
            fill_count = len(df[df['题型'] == 'fill'])
            essay_count = len(df[df['题型'] == 'essay'])
        else:
            single_count = multiple_count = judge_count = fill_count = essay_count = 0
        
        summary_data = {
            '总试题数': total_questions,
            '单选题数': single_count,
            '多选题数': multiple_count,
            '判断题数': judge_count,
            '填空题数': fill_count,
            '简答题数': essay_count
        }
        
        wb, ws = ExcelStyleService.create_styled_worksheet(
            df=df,
            sheet_name='试题数据',
            title='ExamSphere 试题数据导出',
            subtitle=f'导出时间: {pd.Timestamp.now().strftime("%Y-%m-%d %H:%M:%S")}',
            column_widths=column_widths,
            summary_data=summary_data
        )
        
        return wb
    
    @staticmethod
    def create_exam_result_export_excel(df):
        """创建考试结果导出Excel"""
        # 列宽设置
        column_widths = {
            'A': 8,   # ID
            'B': 8,   # 考试ID
            'C': 8,   # 用户ID
            'D': 20,  # 开始时间
            'E': 20,  # 提交时间
            'F': 10,  # 分数
            'G': 10,  # 状态
            'H': 20   # 创建时间
        }
        
        # 汇总数据
        total_records = len(df)
        if total_records > 0:
            avg_score = df['分数'].mean()
            max_score = df['分数'].max()
            min_score = df['分数'].min()
            completed_count = len(df[df['状态'] == 'submitted'])
        else:
            avg_score = max_score = min_score = 0
            completed_count = 0
        
        summary_data = {
            '总记录数': total_records,
            '完成考试数': completed_count,
            '平均分数': f"{avg_score:.2f}",
            '最高分数': max_score,
            '最低分数': min_score
        }
        
        wb, ws = ExcelStyleService.create_styled_worksheet(
            df=df,
            sheet_name='考试结果',
            title='ExamSphere 考试结果导出',
            subtitle=f'导出时间: {pd.Timestamp.now().strftime("%Y-%m-%d %H:%M:%S")}',
            column_widths=column_widths,
            summary_data=summary_data
        )
        
        return wb
    
    @staticmethod
    def create_exam_export_excel(df):
        """创建考试导出Excel"""
        # 列宽设置
        column_widths = {
            'A': 8,   # ID
            'B': 25,  # 考试标题
            'C': 20,  # 科目
            'D': 30,  # 描述
            'E': 10,  # 总分
            'F': 12,  # 题目数量
            'G': 12,  # 考试时长
            'H': 20,  # 开始时间
            'I': 20,  # 结束时间
            'J': 10,  # 状态
            'K': 20   # 创建时间
        }
        
        # 汇总数据
        total_exams = len(df)
        published_count = len(df[df['状态'] == 'published']) if total_exams > 0 else 0
        draft_count = len(df[df['状态'] == 'draft']) if total_exams > 0 else 0
        finished_count = len(df[df['状态'] == 'finished']) if total_exams > 0 else 0
        
        summary_data = {
            '总考试数': total_exams,
            '已发布': published_count,
            '草稿': draft_count,
            '已完成': finished_count
        }
        
        wb, ws = ExcelStyleService.create_styled_worksheet(
            df=df,
            sheet_name='考试数据',
            title='ExamSphere 考试数据导出',
            subtitle=f'导出时间: {pd.Timestamp.now().strftime("%Y-%m-%d %H:%M:%S")}',
            column_widths=column_widths,
            summary_data=summary_data
        )
        
        return wb
